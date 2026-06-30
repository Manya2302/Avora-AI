"""Avora AI — Ensemble Document Processor (Phase 3)"""
import io, re, logging
from dataclasses import dataclass, field
logger = logging.getLogger(__name__)

_PADDLE_OCR_INSTANCE = None

@dataclass
class ExtractionResult:
    raw_text: str = ""
    clean_text: str = ""
    confidence: float = 0.0
    page_count: int = 0
    word_count: int = 0
    engine_used: str = ""
    quality_score: float = 0.0
    language: str = "en"
    tables: list = field(default_factory=list)
    key_value_pairs: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)

class AvoraDocumentProcessor:
    """Ensemble: pdfplumber → PyMuPDF → PaddleOCR → Tesseract. Returns best result."""

    def process(self, file_bytes: bytes, mime_type: str, filename: str = "") -> ExtractionResult:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if mime_type == "application/pdf" or ext == "pdf":
            return self._pdf_ensemble(file_bytes)
        elif ext in ("docx","doc"): return self._docx(file_bytes)
        elif ext in ("xlsx","xls"): return self._excel(file_bytes)
        elif mime_type.startswith("image/") or ext in ("png","jpg","jpeg","tiff"):
            return self._image_ensemble(file_bytes)
        else:
            text = file_bytes.decode("utf-8", errors="replace")
            r = ExtractionResult(raw_text=text, word_count=len(text.split()),
                confidence=1.0, engine_used="plaintext", quality_score=1.0)
            r.clean_text = TextCleaner.clean(text); return r

    def _pdf_ensemble(self, fb):
        results = []
        for fn in (self._pdfplumber, self._pymupdf):
            try:
                r = fn(fb)
                if r.word_count > 10:
                    results.append(r)
            except: pass

        best = None
        if results:
            best = max(results, key=lambda x: (x.quality_score, x.word_count))

        # Fall back to page-by-page OCR if no digital text or digital text is extremely sparse (e.g. scanned PDF)
        if not best or best.word_count < 10:
            logger.info("[Avora OCR] PDF has little or no digital text. Falling back to page-by-page OCR...")
            try:
                import fitz
                doc = fitz.open(stream=fb, filetype="pdf")
                page_texts = []
                total_confidence = 0.0
                pages_processed = 0

                for page_idx, page in enumerate(doc):
                    if page_idx >= 10:  # limit to first 10 pages for performance safety
                        break
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        r_ocr = self._image_ensemble(img_bytes)
                        if r_ocr.raw_text.strip():
                            page_texts.append(r_ocr.raw_text)
                            total_confidence += r_ocr.confidence
                            pages_processed += 1
                    except Exception as e_page:
                        logger.warning(f"[Avora OCR] Error OCRing page {page_idx}: {e_page}")

                if page_texts:
                    ocr_text = "\n\n".join(page_texts)
                    avg_conf = total_confidence / pages_processed if pages_processed > 0 else 0.8
                    best = ExtractionResult(
                        raw_text=ocr_text,
                        page_count=len(doc),
                        word_count=len(ocr_text.split()),
                        confidence=avg_conf,
                        engine_used="pdf_page_ocr",
                        quality_score=QualityScorer.score(ocr_text, avg_conf)
                    )
            except Exception as ocr_err:
                logger.error(f"[Avora OCR] PDF page OCR fallback failed: {ocr_err}")

        if not best:
            return ExtractionResult(warnings=["No text extracted"])

        best.clean_text = TextCleaner.clean(best.raw_text)
        best.key_value_pairs = KeyValueExtractor.extract(best.clean_text)
        return best

    def _pdfplumber(self, fb):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(fb)) as pdf:
            text = "\n\n".join(p.extract_text() or "" for p in pdf.pages)
            wc   = len(text.split())
            conf = 0.96 if wc > 100 else 0.70
            return ExtractionResult(raw_text=text, page_count=len(pdf.pages),
                word_count=wc, confidence=conf, engine_used="pdfplumber",
                quality_score=QualityScorer.score(text, conf))

    def _pymupdf(self, fb):
        import fitz
        doc  = fitz.open(stream=fb, filetype="pdf")
        text = "\n\n".join(p.get_text() for p in doc)
        wc   = len(text.split())
        conf = 0.95 if wc > 50 else 0.60
        return ExtractionResult(raw_text=text, page_count=len(doc),
            word_count=wc, confidence=conf, engine_used="pymupdf",
            quality_score=QualityScorer.score(text, conf))

    def _image_ensemble(self, fb):
        r = self._paddleocr(fb)
        if r.quality_score >= 0.7:
            r.clean_text = TextCleaner.clean(r.raw_text)
            r.key_value_pairs = KeyValueExtractor.extract(r.clean_text)
            return r
        r2 = self._tesseract(fb)
        best = r if r.quality_score >= r2.quality_score else r2
        best.clean_text = TextCleaner.clean(best.raw_text)
        best.key_value_pairs = KeyValueExtractor.extract(best.clean_text)
        return best

    def _paddleocr(self, fb):
        global _PADDLE_OCR_INSTANCE
        try:
            import numpy as np, cv2
            arr    = np.frombuffer(fb, np.uint8)
            img    = cv2.imdecode(arr, cv2.IMREAD_COLOR)
            img    = ImagePreprocessor.enhance(img)
            
            # Cache PaddleOCR instance to avoid loading weights on every page/call
            if '_PADDLE_OCR_INSTANCE' not in globals() or globals()['_PADDLE_OCR_INSTANCE'] is None:
                from paddleocr import PaddleOCR
                globals()['_PADDLE_OCR_INSTANCE'] = PaddleOCR(use_angle_cls=True, lang="en", show_log=False, use_gpu=False)
            
            ocr = globals()['_PADDLE_OCR_INSTANCE']
            result = ocr.ocr(img, cls=True)
            if result and result[0]:
                lines = [l[1][0] for l in result[0] if l]
                confs = [l[1][1] for l in result[0] if l]
                text  = "\n".join(lines)
                avg   = sum(confs)/len(confs) if confs else 0.8
                return ExtractionResult(raw_text=text, word_count=len(text.split()),
                    confidence=round(avg,3), engine_used="paddleocr",
                    quality_score=QualityScorer.score(text,avg))
        except Exception as e:
            logger.debug(f"[PaddleOCR] {e}")
        return ExtractionResult(engine_used="paddleocr_failed")

    def _tesseract(self, fb):
        try:
            import pytesseract
            import numpy as np, cv2
            from PIL import Image
            arr   = np.frombuffer(fb, np.uint8)
            img   = cv2.imdecode(arr, cv2.IMREAD_COLOR)
            img   = ImagePreprocessor.enhance(img)
            pil   = Image.fromarray(cv2.cvtColor(img, cv2.COLOR_BGR2RGB))
            data  = pytesseract.image_to_data(pil, output_type=pytesseract.Output.DICT)
            words = [w for w,c in zip(data["text"],data["conf"]) if w.strip() and int(c)>0]
            confs = [int(c)/100 for c in data["conf"] if int(c)>0]
            text  = " ".join(words)
            avg   = sum(confs)/len(confs) if confs else 0.0
            return ExtractionResult(raw_text=text, word_count=len(words),
                confidence=round(avg,3), engine_used="tesseract",
                quality_score=QualityScorer.score(text,avg))
        except Exception as e:
            logger.debug(f"[Tesseract] {e}")
        return ExtractionResult(engine_used="tesseract_failed")

    def _doc_fallback(self, fb):
        # A simple fallback for legacy .doc files to extract text chunks using regex
        try:
            import re
            # Extract printable strings from binary data
            text_parts = re.findall(r"[\x20-\x7E\u00A0-\uFFFF]{4,}", fb.decode("utf-8", errors="ignore"))
            clean_parts = []
            for p in text_parts:
                p_clean = p.strip()
                # Filter out junk
                if len(p_clean) > 8 and any(c.isalpha() for c in p_clean):
                    clean_parts.append(p_clean)
            text = "\n".join(clean_parts)
            r = ExtractionResult(
                raw_text=text,
                word_count=len(text.split()),
                confidence=0.5,
                engine_used="doc_binary_extractor",
                quality_score=QualityScorer.score(text, 0.5)
            )
            r.clean_text = TextCleaner.clean(text)
            r.key_value_pairs = KeyValueExtractor.extract(r.clean_text)
            return r
        except Exception as e:
            return ExtractionResult(engine_used="doc_fallback_failed", warnings=[str(e)])

    def _docx(self, fb):
        try:
            from docx import Document as D
            doc   = D(io.BytesIO(fb))
            text  = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            r     = ExtractionResult(raw_text=text, word_count=len(text.split()),
                confidence=0.99, engine_used="python-docx",
                quality_score=QualityScorer.score(text,0.99))
            r.clean_text = TextCleaner.clean(text)
            r.key_value_pairs = KeyValueExtractor.extract(r.clean_text)
            return r
        except Exception as e:
            logger.info(f"[Avora Doc] python-docx failed, falling back to legacy .doc binary extractor: {e}")
            return self._doc_fallback(fb)

    def _excel(self, fb):
        try:
            import openpyxl
            wb    = openpyxl.load_workbook(io.BytesIO(fb), data_only=True)
            cells = [str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None]
            text  = " ".join(cells)
            r     = ExtractionResult(raw_text=text, word_count=len(cells),
                confidence=0.99, engine_used="openpyxl",
                quality_score=QualityScorer.score(text,0.99))
            r.clean_text = TextCleaner.clean(text); return r
        except Exception as e:
            return ExtractionResult(engine_used="excel_failed", warnings=[str(e)])


class ImagePreprocessor:
    @staticmethod
    def enhance(img):
        try:
            import cv2
            gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
            t    = cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY,11,2)
            d    = cv2.fastNlMeansDenoising(t, h=10)
            h, w = d.shape[:2]
            if w < 1000: d = cv2.resize(d,(w*2,h*2),interpolation=cv2.INTER_CUBIC)
            return d
        except: return img


class TextCleaner:
    @staticmethod
    def clean(text: str) -> str:
        if not text: return ""
        text = re.sub(r"\r\n|\r","\n",text)
        text = re.sub(r"\n{3,}","\n\n",text)
        text = re.sub(r" {2,}"," ",text)
        text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E\u00A0-\uFFFF]"," ",text)
        return "\n".join(l.strip() for l in text.split("\n") if len(l.strip())>1).strip()


class QualityScorer:
    @staticmethod
    def score(text: str, conf: float) -> float:
        if not text.strip(): return 0.0
        words = text.split()
        if not words: return 0.0
        alpha = sum(1 for w in words if re.match(r"^[a-zA-Z]",w)) / len(words)
        length= min(len(words)/200, 1.0)
        garb  = len(re.findall(r"[^a-zA-Z0-9\s.,;:!?₹$%&()\-\'\"/]",text)) / max(len(text),1)
        return round(alpha*0.3 + length*0.2 + conf*0.3 + (1-min(garb*5,1))*0.2, 3)


class KeyValueExtractor:
    PATTERNS = [
        (r"(?:invoice|bill)\s*(?:no|number|#)[:\s]+([A-Z0-9\-/]+)","invoice_number"),
        (r"(?:gstin|gst number)[:\s]+(\d{2}[A-Z]{5}\d{4}[A-Z]\d[Z][A-Z\d])","gstin"),
        (r"(?:pan)[:\s]+([A-Z]{5}\d{4}[A-Z])","pan"),
        (r"(?:date|dated)[:\s]+(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})","date"),
        (r"(?:total|amount|grand total)[:\s]+(?:₹|Rs\.?|INR)?\s*([\d,]+(?:\.\d{2})?)","total_amount"),
        (r"(?:due date|payment due)[:\s]+(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})","due_date"),
        (r"(?:vendor|supplier|from)[:\s]+([A-Z][^\n]{5,60})","vendor"),
    ]
    @staticmethod
    def extract(text: str) -> dict:
        result = {}
        for pat, key in KeyValueExtractor.PATTERNS:
            m = re.search(pat, text, re.I)
            if m: result[key] = m.group(1).strip()
        return result
